from datetime import datetime
from pathlib import Path
import hashlib
import json

from taiyi_agent.memory.rag.models import SourceDocument, TextChunk


def load_file(path: str, encoding: str = "utf-8") -> SourceDocument:
    """
    加载不同类型的文件

    Args：
        path：文件路径
    Return：
        SourceDocument： RAG系统原始文档类
    """
    file = Path(path)
    suffix = file.suffix.lower()

    if suffix in {".txt", ".md", ".py"}:
        content = file.read_text(encoding=encoding)
        source_type = "text"

    elif suffix == ".json":
        content = extract_json_text(file, encoding)
        source_type = "json"

    elif suffix == ".pdf":
        content = extract_pdf_text(file)
        source_type = "pdf"

    elif suffix == ".docx":
        content = extract_docx_text(file)
        source_type = "docx"

    elif suffix == ".csv":
        content = extract_csv_text(file, encoding)
        source_type = "csv"

    else:
        raise ValueError(f"暂不支持的文件类型: {suffix}")

    content_hash = hashlib.sha256(content.encode(encoding)).hexdigest()
    doc_id = hashlib.sha256(
        f"{file.resolve()}\0{content_hash}".encode(encoding)
    ).hexdigest()

    print(f"[load_file], {source_type}加载成功, doc_id:{doc_id}, 文档部分内容:{content[:30]}")

    return SourceDocument(
        doc_id=doc_id,
        content=content,
        metadata={
            "source_uri": str(file.resolve()),
            "source_type": source_type,
            "title": file.stem,
        },
    )

def extract_json_text(file: Path, encoding: str) -> str:
    """
    提取 json 文件文本内容
    """
    data = json.loads(file.read_text(encoding=encoding))
    return json.dumps(data, ensure_ascii=False, indent=2)

def extract_pdf_text(file: Path) -> str:
    """
    提取 pdf 文件文本内容
    """
    from pypdf import PdfReader

    reader = PdfReader(str(file))
    return "\n\n".join(
        f"[page={i + 1}]\n{page.extract_text() or ''}"
        for i, page in enumerate(reader.pages)
    )

def extract_docx_text(file: Path) -> str:
    """
    提取 docx 文件文本内容
    """
    from docx import Document
    doc = Document(str(file))
    parts = []

    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    # 读取表格
    for table in doc.tables:
        for row in table.rows:
            row_text = "\t".join(cell.text.strip() for cell in row.cells)
            if row_text:
                parts.append(row_text)
    return "\n".join(parts)

def extract_csv_text(file: Path, encoding: str) -> str:
    """
    提取 csv 文件文本内容
    """
    import csv
    rows_out = []

    with open(file, encoding=encoding, newline="") as f:
        rows = csv.reader(f)
        for idx, row in enumerate(rows):
            # 转义单元格内的 | 符号
            escaped = [cell.replace("|", r"\|") for cell in row]
            rows_out.append("| " + " | ".join(escaped) + " |")
            # 表头之后插入markdown分割行
            if idx == 0:
                rows_out.append("|" + "---|" * len(row))
    return "\n".join(rows_out)
